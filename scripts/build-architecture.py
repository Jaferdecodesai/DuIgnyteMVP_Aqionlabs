"""Refresh the supplied system design for the actual Cloudflare MVP.

Retains its twelve-section technical narrative; emits editable Word and a
readable repository companion. Requires python-docx and Pillow.
"""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
CFG = json.loads((ROOT/'pitch/build-config.json').read_text())
MVP, GH = CFG['mvp'], CFG['github']
OUT = ROOT/'architecture'
OUT.mkdir(exist_ok=True)
NAVY, TEAL, PALE = '102F3A', '087F80', 'EDF4F2'
d = Document()
s = d.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin, s.bottom_margin = Inches(.68), Inches(.62)
s.left_margin, s.right_margin = Inches(.8), Inches(.8)
normal = d.styles['Normal']
normal.font.name, normal.font.size = 'Calibri', Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08
for name, size in [('Title',27),('Subtitle',15),('Heading 1',15),('Heading 2',11)]:
 st=d.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(7)
 if name.startswith('Heading'):st.font.bold=True;st.paragraph_format.keep_with_next=True
for st in d.styles:
 for border in st.element.xpath('.//w:pBdr'):
  border.getparent().remove(border)
header=s.header.paragraphs[0]
header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=header.add_run('aqionlabs.ai');r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0)
footer=s.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Aqion Relay | System design | ');r.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
md=['# Aqion Relay system design','',f'Updated 16 October 2026. Brand: aqionlabs.ai.','',f'- [Live MVP]({MVP})',f'- [GitHub repository]({GH})',f'- [Editable architecture DOCX](Aqion-Relay-System-Design.docx)',f'- [Architecture PDF](Aqion-Relay-System-Design.pdf)',f'- [Pitch deck](Aqion-Relay-Du-Ignyte-Pitch.pdf)','']

def p(text='',style=None):
 paragraph=d.add_paragraph(text,style)
 if text:md.extend([text,''])
 return paragraph
def h(text):
 d.add_heading(text,1);md.extend(['## '+text,''])
def page():d.add_page_break()
def link(par,label,url):
 rel=par.part.relate_to(url,RT.HYPERLINK,is_external=True)
 x=OxmlElement('w:hyperlink');x.set(qn('r:id'),rel)
 run=OxmlElement('w:r');props=OxmlElement('w:rPr')
 c=OxmlElement('w:color');c.set(qn('w:val'),TEAL);props.append(c)
 u=OxmlElement('w:u');u.set(qn('w:val'),'single');props.append(u)
 run.append(props);t=OxmlElement('w:t');t.text=label;run.append(t);x.append(run);par._p.append(x)
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,width in enumerate(widths):t.columns[i].width=Inches(width)
 for i,label in enumerate(headers):t.rows[0].cells[i].text=label
 for row in rows:
  cells=t.add_row().cells
  for i,val in enumerate(row):cells[i].text=str(val)
 for ri,row in enumerate(t.rows):
  for ci,cell in enumerate(row.cells):
   cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcPr=cell._tc.get_or_add_tcPr()
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   tcPr.append(borders)
   margins=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');margins.append(el)
   tcPr.append(margins)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),NAVY if ri==0 else ('F2F6F5' if ri%2 else 'FFFFFF'));tcPr.append(shade)
   for par in cell.paragraphs:
    par.paragraph_format.space_after=Pt(2);par.paragraph_format.line_spacing=1.03
    for run in par.runs:
     run.font.size=Pt(10);run.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000');run.bold=ri==0
   if ri==0:
    repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
   no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
 md.extend(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |'])
 md.extend('| '+' | '.join(str(v).replace('\n',' ') for v in row)+' |' for row in rows);md.append('')
 d.add_paragraph().paragraph_format.space_after=Pt(2)
 return t

# Technical diagram is drawn from geometry, not copied from the outdated image.
im=Image.new('RGB',(1600,700),'#F2F7F5');draw=ImageDraw.Draw(im)
fontroot=ROOT/'pitch/assets'
f=lambda size,bold=False:ImageFont.truetype(str(fontroot/('InstrumentSans-Bold.ttf' if bold else 'InstrumentSans-Regular.ttf')),size)
draw.rectangle((0,0,1600,92),fill='#'+NAVY)
draw.text((40,22),'Cloudflare MVP  |  Current execution boundary',font=f(29,True),fill='white')
boxes=[(45,'Browser','Next.js / React','Signed session + CSRF'),(565,'Worker','Static assets + routing','Canonical HTTPS origin'),(1085,'Durable Object','API + engine + SQLite','Atomic draft + audit')]
for x,title,sub,body in boxes:
 draw.rounded_rectangle((x,150,x+465,345),radius=12,fill='white',outline='#'+TEAL,width=3)
 draw.text((x+25,177),title,font=f(31,True),fill='#'+NAVY)
 draw.text((x+25,239),sub,font=f(24),fill='#'+NAVY)
 draw.text((x+25,291),body,font=f(22),fill='#'+TEAL)
for x in [510,1030]:
 draw.line((x,245,x+48,245),fill='#'+TEAL,width=5);draw.polygon([(x+48,245),(x+33,236),(x+33,254)],fill='#'+TEAL)
draw.line((1320,345,1320,400),fill='#'+TEAL,width=4)
draw.rounded_rectangle((1085,400,1550,520),radius=10,fill='#D3EDE0')
draw.text((1110,421),'Snapshot / proposal / draft / audit',font=f(22,True),fill='#'+NAVY)
draw.text((1110,466),'transactionSync  |  1.5 MiB cap',font=f(22),fill='#'+NAVY)
draw.text((45,420),'CURRENT: synthetic records and unsent drafts',font=f(25,True),fill='#'+NAVY)
draw.text((45,468),'One coordinator; 200 anonymous workspaces; seven-day retention.',font=f(23),fill='#'+NAVY)
draw.line((45,570,1550,570),fill='#9DBDB5',width=2)
draw.text((45,596),'PILOT GATES: tenant IAM  /  scoped ERP adapter  /  outbox + reconciliation',font=f(24,True),fill='#'+TEAL)
draw.text((45,644),'du identity and distribution, processing location, restore and load acceptance require sponsor approval.',font=f(22),fill='#'+NAVY)
im.save(OUT/'current-architecture.png')

p('Aqion Relay System Design','Title')
p('Cloudflare MVP and pilot integration','Subtitle')
p('16 October 2026  |  ResilienceTech  |  aqionlabs.ai')
p('This document updates the supplied system design to the current mobile web implementation. It describes the deployed execution boundary, API contracts, approval integrity and release operations, then identifies the integration and scale gates for a proposed du SME pilot.')
table(['Status','Current scope'],[
 ('MVP implementation','Next.js frontend, native Cloudflare Worker API and SQLite-backed Durable Object'),
 ('Data and actions','Fictional SME records; owner-approved unsent drafts; no supplier or payment transmission'),
 ('Product and engineering','aqionlabs.ai; sponsor security, integration and pilot roles to be appointed'),
 ('Pilot proposal','160 consenting SMEs in four waves; minimum 150 enrolled before intervention'),
],[1.5,5.4])
for label,url in [('Live MVP',MVP),('GitHub repository',GH),('Pitch deck',MVP+'/pitch/'),('Architecture document',CFG['architecture'])]:
 par=d.add_paragraph();par.add_run(label+'\n').bold=True;link(par,url,url)
p('The repository README provides direct access to the live MVP, pitch formats and architecture formats. Current release evidence is kept with the source; targets and planned integrations are distinguished from verified implementation.')

page();h('1 Abstract')
p('Aqion Relay is a ResilienceTech decision workbench for UAE non-perishable goods distributors. It compares supply recovery purchases against an owner-set operating cash floor and fictional approved-supplier constraints. The implemented web MVP uses a numerical forecast, a 30-day stock and scheduled-cash simulation, and explicit owner review before saving an unsent draft.')
p('The public deployment serves a Next.js static export through Cloudflare Worker assets. Its API, engine and domain service execute inside a native SQLite-backed Durable Object. The prototype scope remains one SKU per SME and three fictional SMEs. ERP connections, generative narration, authenticated business tenants and du distribution enter through the pilot roadmap.')
h('2 Goals and non goals')
table(['Goals','Non goals'],[
 ('Compare recovery choices with daily cash and service impact','Guarantee actual cash, supplier delivery or a global optimum'),
 ('Keep source-bound human review and durable audit evidence','Send purchases, move funds or replace an ERP'),
 ('Test 40% shorter planning time in controlled pilot drills','Claim achieved customer benefits or completed recruitment'),
 ('Provide digital integration and a gated 40,000-SME scale path','Assert the capped demo already supports enterprise scale'),
],[3.45,3.45])
h('3 Background and problem statement')
p('A shipment delay threatens customer fulfilment while a replacement purchase competes with payroll, rent and other cash obligations. Relay focuses on that decision alongside existing inventory tools: read a versioned operational snapshot, stress-test a disruption, compare feasible alternatives, and retain the reviewed evidence.')
p('The reference synthetic packaging case evaluates 130 combinations. A 400-unit Oasis recovery purchase costs AED 7,887.50 and serves 1,362 projected units versus 990 without recovery. Minimum projected cash is AED 28,112.50 against an AED 24,000 floor. These are reproducible fictional model outputs, not customer outcomes.')

page();h('4 Current architecture and pilot extension')
d.add_picture(str(OUT/'current-architecture.png'),width=Inches(6.9))
par=d.add_paragraph('Figure 1. Current Cloudflare boundary and the separately gated pilot extension.');par.paragraph_format.space_after=Pt(9)
md.extend(['![Current Cloudflare architecture](architecture/current-architecture.png)',''])
table(['Component','Responsibility','Persistence and failure behavior'],[
 ('Next.js and React','Responsive Overview, Evidence, Connections and Pilot views','Static assets; retry UI on API failure; no client-selected workspace key'),
 ('Worker router','Canonical origin, security headers, static assets and API allowlist','Forward intact Request to one stable coordinator; reject unknown API paths'),
 ('Durable Object API','Signed anonymous session, input validation and domain actions','Native SQLite; transactionSync encloses read, checks, mutation, encode and commit'),
 ('Forecast and simulator','Ridge weekday/trend versus seasonal naive; bounded candidate enumeration','Label sparse-history fallback; maintain baseline and cash-floor checks'),
 ('Approval service','Recheck source/calculation versions and explicit confirmation','One outstanding draft per SME; atomic draft and audit; idempotent repeat'),
 ('Pilot connectors','Consented ERP reads, approved draft writes and du interfaces','Proposed scoped adapters, outbox, reconciliation and tenant IAM'),
],[1.35,2.6,2.95])
p('The singleton uses idFromName("relay-demo-v1"). Demo bounds are 200 anonymous workspaces, seven-day expiry, 1.5 MiB per document, 100 undrafted proposals plus referenced drafts, and 200 drafts/audit events per SME. Expired workspace cleanup is opportunistic during creation. Pilot scale requires partitioned authenticated tenant/SME storage and separate acceptance tests.')

page();h('5 Request lifecycle')
for text in [
 '1. GET /api/state issues or verifies a signed anonymous session and returns the selected fictional SME plus a CSRF token.',
 '2. The browser POSTs JSON with the canonical Origin, session cookie and X-Relay-CSRF. Inputs are bounded to 256 KiB.',
 '3. The engine validates source age and supplier terms, fits a forecast, and evaluates 30-day stock and scheduled cash under the scenario.',
 '4. The owner compares alternatives, opens the review dialog, checks quantities and cash floor, then explicitly confirms an unsent draft.',
 '5. Approval rechecks the current source digest, calculation hash, proposal expiry, plan eligibility and outstanding-draft invariant in one synchronous SQL transaction.',
 '6. The service persists the draft and its audit entry together. An identical retry returns the same draft, including canceled status.',
 '7. Evidence can be exported, a draft canceled, or a confirmed fictional snapshot reset. Imports invalidate prior source-bound proposals.',
]:p(text)
h('6 API and data contracts')
table(['Method and route','Contract'],[
 ('GET state / export','tenant=packaging, trading or homeware; signed browser session; export contains no session or CSRF secret'),
 ('POST simulate','tenant and bounded delay, demand, freight and cash_floor; returns source/version-bound proposal and plans'),
 ('POST approve / cancel','proposal_id, plan_id and confirm=true for approval; draft_id for cancellation'),
 ('POST import / reset','Complete tenant-matched snapshot with synthetic=true; reset requires confirm=true'),
 ('GET health','Native storage readiness, synthetic mode and writes flag; no credentials or private identifiers'),
 ('POST access / logout','Optional shared demo gate and cookie clearing; gate does not provide enterprise identity'),
],[2.1,4.8])
p('The /api/ prefix is explicit: unknown or nested actions return JSON 404 and unsupported methods return 405. The authoritative contract is docs/WEB_API.md; TypeScript shapes are in src/lib/types.ts. A snapshot includes stock, scheduled receipts/obligations, history, prepaid inbound stock, approved suppliers, quote validity and modeled age. Proposal lifetime is 30 minutes.')

page();h('7 Consistency idempotency and replay')
p('Native SQLite transactionSync performs the workspace read, domain checks, mutation, document-size check and persisted write synchronously. A thrown validation or storage error rolls back both draft and audit. The transaction prevents simultaneous requests from creating two outstanding drafts. An identical approval returns its existing result; a canceled draft is not resurrected.')
p('Proposal identity binds the snapshot digest, scenario and calculation versions. An import changes the source revision, and approval recomputes the result before accepting the reviewed plan. Stale source data, expired supplier terms, elapsed proposal lifetime, an infeasible or fabricated plan, and an existing outstanding draft reject approval.')
p('These guarantees apply to the unsent demo draft. An enterprise write requires a reservation and transactional outbox, a provider idempotency strategy, event/revision tracking and reconciliation. Ambiguous ERP results remain pending reconciliation; a blind retry must not create a second external order.')
h('8 Security and privacy considerations')
for text in [
 'Current sessions: HMAC-signed anonymous browser cookie; HttpOnly, SameSite=Strict, Secure on HTTPS; seven-day lifetime. This isolates fictional workspaces but is not tenant IAM.',
 'Current request controls: exact canonical HTTPS origin/Host, same-origin mutation checks, HMAC CSRF, bounded JSON, strict validation and safe error messages. Security headers apply to assets and API responses.',
 'Current secrets: the stable signing secret is a Cloudflare Worker secret. Local .env, .dev.vars, storage, OAuth credentials and persistence-probe cookies are excluded from public source and static asset uploads.',
 'Current writes: RELAY_WRITES_ENABLED controls demo mutations. Enabled actions save fictional state only. Per-session rate limits are best effort per running coordinator instance, not a distributed attack-control claim.',
 'Pilot access: authenticate business users, authorize every tenant object/job/export, use scoped connector credentials and redact telemetry. Keep operational records out of cross-SME prompts and analysis.',
 'Pilot processing: agree data roles, location, retention and deletion; test encryption, restore, incident ownership and audited write-pausing. Global edge deployment does not by itself establish UAE data residency.',
]:p(text)

page();h('9 Operational readiness and release controls')
p('The release process is source review, locked dependency install, strict type checks, regression tests and coverage, Node build, Next.js static export, Worker dry run, native local smoke, owner-authorized deployment and public verification. GitHub Actions runs build/test/audit gates on pushes. Cloudflare production deployment remains an explicit release action; no stored CI cloud credential is introduced.')
table(['Gate','Acceptance evidence'],[
 ('Code and dependencies','127 regression tests and configured coverage thresholds; production/static builds; zero moderate-or-higher audit findings'),
 ('Native runtime','Actual workerd and HTTPS API checks for sessions, routing, calculations, CSRF, isolation, imports and concurrent approvals'),
 ('Persistence and assets','Retain draft plus audit across same-namespace redeploy; downloaded presentation/document bytes match release SHA-256'),
 ('Browser and documents','Scenario and owner review; responsive layout; links and all final slides/pages rendered and inspected'),
 ('Pilot service targets','Proposed 99.9% availability, p95 scenario <5 s, RPO <=15 min and RTO <=4 h; sponsor agreement and separate acceptance required'),
],[1.65,5.25])
p('Routine updates retain Worker name, Durable Object class/namespace, stable idFromName value and signing secret. The requested hostname is a new named deployment and starts a separate anonymous demo workspace. To pause mutations set RELAY_WRITES_ENABLED=false and redeploy. For a failed code release select a verified prior version with wrangler rollback; code rollback does not undo persisted state. Never delete the namespace as a routine update.')
h('10 Alternatives considered')
table(['Alternative','Decision'],[
 ('Static frontend only','Rejected: server authority, durable approvals and audit must execute on the server'),
 ('Node plus MySQL hosting','Optional adapter retained; Cloudflare native SQLite is the deployed target'),
 ('Full micro ERP or liquidity marketplace','Deferred: migration or counterparty/legal scope exceeds the first recovery decision'),
 ('LLM decides quantities or autonomous buying','Rejected for the MVP: numerical constraints and owner authority control each draft'),
],[2.15,4.75])

page();h('11 Integration decisions before the pilot')
p('Confirm the du sponsor and technical owners, consenting SME recruitment, supported ERP organization/region and OAuth scopes, processing/privacy requirements, support ownership and cost ceiling. Screen source completeness and approved alternative supplier terms before admitting an SME to intervention. No du, ERP, banking or supplier interface is claimed connected.')
h('12 Integration roadmap and pilot execution')
table(['Phase','Delivery and acceptance gate'],[
 ('Weeks 1 to 2','Map consented snapshots, source IDs and supplier terms; approve scope and processing/security design'),
 ('Weeks 3 to 4','Zoho read-only adapter; normalized stock, sales and scheduled-cash mapping; reconcile source totals and revisions'),
 ('Weeks 5 to 6','Owner-authorized ERP drafts, outbox and idempotency; sandbox ambiguity, retry, authorization and restore tests'),
 ('Weeks 7 to 8','Agreed du interfaces, monitoring and support; named sponsor acceptance before pilot comparison'),
],[1.45,5.45])
p('Proposed enrollment is 160 consenting non-perishable-goods SMEs, with at least 150 enrolled before intervention. Four waves of 40 begin in pilot weeks 1, 2, 3 and 4. Each receives two baseline weeks and six comparison weeks. Within each wave and sector/size/ERP strata, randomize 20 immediate-access and 20 delayed-access SMEs. The last comparison ends in week 11 and analysis occurs in week 12. Do not replace randomized dropouts.')
p('Prospective targets are 40% shorter time to an owner-approved feasible plan, 90% rubric-passing plans across assigned drills, and 70% weekly workflow use among SMEs with access. Pre-register the analysis, complete power/MDE work before intervention, and report effects, uncertainty, attrition and deviations. Stop on a projected floor breach, unauthorized external write or duplicate external effect.')
p('Scale from the accepted pilot to a controlled 1,000-SME rollout, then the 40,000+ ambition only after isolation, peak-load, unit-cost and recovery acceptance. Use authenticated tenant/SME partitioning, queues, shared limits and observability. The 200-workspace singleton demo is not evidence of that capacity.')
par=d.add_paragraph();par.add_run('Implementation and evidence: ').bold=True
link(par,'repository README',GH+'#readme');par.add_run(' | ');link(par,'API contract',GH+'/blob/main/docs/WEB_API.md');par.add_run(' | ');link(par,'release runbook',GH+'/blob/main/docs/CLOUDFLARE.md')
par=d.add_paragraph();par.add_run('Primary platform references: ')
link(par,'Worker static assets','https://developers.cloudflare.com/workers/static-assets/');par.add_run(' | ');link(par,'Durable Object SQLite','https://developers.cloudflare.com/durable-objects/api/sqlite-storage-api/')
p('This design records implemented behavior and a proposed pilot. Portal registration, final upload and a submission receipt remain owner actions; no submitted entry or organizer approval is asserted.')
d.core_properties.title='Aqion Relay System Design'
d.core_properties.author='aqionlabs.ai'
d.core_properties.subject='Current Cloudflare MVP and du SME pilot integration'
d.save(ROOT/'Aqion-Relay-System-Design.docx')
(ROOT/'ARCHITECTURE.md').write_text('\n'.join(md).rstrip()+'\n')
print(json.dumps({'docx':str(ROOT/'Aqion-Relay-System-Design.docx'),'sections':12,'tables':len(d.tables),'diagram':str(OUT/'current-architecture.png')}))
