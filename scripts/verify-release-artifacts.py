"""Check the requested cross-format release links and editable source integrity."""
from pathlib import Path
from zipfile import ZipFile
import json, hashlib, xml.etree.ElementTree as ET
from pypdf import PdfReader
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
cfg=json.loads((ROOT/'pitch/build-config.json').read_text())
scene=json.loads((ROOT/'pitch/scene.json').read_text())
ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
    'p':'http://schemas.openxmlformats.org/presentationml/2006/main',
    'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
expected={1:{cfg['mvp'],cfg['github']},7:{cfg['mvp'],cfg['github']},
          8:{cfg['architecture']},14:{cfg['mvp'],cfg['github']}}
pdf=PdfReader(ROOT/'Aqion-Relay-Du-Ignyte-Pitch.pdf')
assert len(pdf.pages)==14
for page in pdf.pages:assert 'aqionlabs.ai' in page.extract_text()
for index,urls in expected.items():
 links={a.get_object().get('/A',{}).get('/URI') for a in pdf.pages[index-1].get('/Annots',[])}
 assert urls<=links,(index,urls-links)
cover=[e for e in scene['slides'][0]['elements'] if e.get('url') in expected[1]]
assert len(cover)==2 and cover[0]['x']==cover[1]['x']==310
assert all('\n' not in e['text'] for e in cover)
with ZipFile(ROOT/'Aqion-Relay-Du-Ignyte-Pitch.pptx') as z:
 assert z.testzip() is None
 for index,urls in expected.items():
  relations=ET.fromstring(z.read(f'ppt/slides/_rels/slide{index}.xml.rels'))
  targets={r.attrib['Id']:r.attrib['Target'] for r in relations if 'hyperlink' in r.attrib.get('Type','')}
  assert urls<=set(targets.values()),index
  xml=ET.fromstring(z.read(f'ppt/slides/slide{index}.xml'))
  if index==7:
   image_links=xml.findall('.//p:pic/p:nvPicPr/p:cNvPr/a:hlinkClick',ns)
   assert any(targets.get(e.attrib.get('{'+ns['r']+'}id'))==cfg['mvp'] for e in image_links)
doc=Document(ROOT/'Aqion-Relay-System-Design.docx')
with ZipFile(ROOT/'Aqion-Relay-System-Design.docx') as z:
 assert z.testzip() is None
 rels=ET.fromstring(z.read('word/_rels/document.xml.rels'))
 assert {cfg['mvp'],cfg['github']}<={r.attrib.get('Target') for r in rels}
arch_pdf=PdfReader(ROOT/'Aqion-Relay-System-Design.pdf')
assert len(arch_pdf.pages)==7
text='\n'.join(p.text for p in doc.paragraphs)
for value in ['transactionSync','1.5 MiB','200 anonymous','256 KiB','40,000','127 regression']:
 assert value in text or any(value in c.text for t in doc.tables for row in t.rows for c in row.cells),value
outputs=['Aqion-Relay-Du-Ignyte-Pitch.'+e for e in ['pdf','pptx','html']]+['Aqion-Relay-System-Design.'+e for e in ['docx','pdf']]
files=[]
for name in outputs:
 data=(ROOT/name).read_bytes();assert len(data)<40_000_000
 files.append({'file':name,'bytes':len(data),'sha256':hashlib.sha256(data).hexdigest()})
for ext in ['pdf','pptx','html']:
 assert (ROOT/f'Aqion-Relay-Du-Ignyte-Pitch.{ext}').read_bytes()==(ROOT/f'pitch/Aqion-Relay-Du-Ignyte-Pitch.{ext}').read_bytes()
report={'pitch_slides':14,'architecture_pages':7,'cover_url_column':310,
        'brand_on_all_slides':'aqionlabs.ai','links_verified_on_slides':[1,7,8,14],
        'dashboard_image_hyperlinked':True,'files':files}
(ROOT/'evidence/ARTIFACT_LINK_QA.json').write_text(json.dumps(report,indent=2)+'\n')
print(json.dumps(report,indent=2))
